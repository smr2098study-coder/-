"""Regression tests for strict controls and content-preserving cleanup."""
import importlib.util
import json
from pathlib import Path
import tempfile
import unittest
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree as E

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / 'plugins/shimen-format/skills/shimen-word-format/scripts/word_format.py'
spec = importlib.util.spec_from_file_location('word_format', SCRIPT)
fmt = importlib.util.module_from_spec(spec)
spec.loader.exec_module(fmt)


class WordFormatTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def build_sample(self):
        from PIL import Image
        Image.new('RGB', (400, 180), 'white').save(self.path / 'figure.png')
        blocks = []
        for level, text in enumerate(['一、背景', '（一）目的', '1、方法', '（1）步骤', '①范围'], 1):
            blocks.append({'type': 'heading', 'level': level, 'text': text})
            blocks.append({'type': 'paragraph', 'text': '测试正文 English 2026，验证字体与段落格式。'})
        blocks.extend([
            {'type': 'page_break'},
            {'type': 'table', 'caption': '表1 汇总', 'rows': [['指标', 'Value'], ['样本', '20']]},
            {'type': 'image', 'path': 'figure.png', 'caption': '图1 流程'}])
        source = self.path / 'input.json'
        source.write_text(json.dumps({'title': '格式验收', 'blocks': blocks}, ensure_ascii=False), encoding='utf-8')
        target = self.path / 'sample.docx'
        fmt.build(source, target)
        return target

    def test_all_profiles_and_image_caption(self):
        path = self.build_sample()
        result = fmt.audit(path)
        self.assertEqual(result['errors'], 0, result['issues'])
        self.assertEqual(result['warnings'], 0, result['issues'])
        pkg = fmt.Package(path)
        for p in pkg.paragraphs:
            if pkg.role(p) == 'body':
                self.assertFalse(fmt.enabled(pkg.value(p, 'keepNext')))
            if pkg.role(p).startswith('h'):
                self.assertFalse(fmt.enabled(pkg.value(p, 'keepNext')))
        images = [p for p in pkg.paragraphs if pkg.role(p) == 'image']
        self.assertEqual(len(images), 1)
        self.assertFalse(fmt.enabled(pkg.value(images[0], 'keepNext')))
        captions = [p for p in pkg.paragraphs if pkg.role(p) == 'caption']
        self.assertFalse(fmt.enabled(pkg.value(captions[0], 'keepNext')))
        self.assertFalse(fmt.enabled(pkg.value(captions[1], 'keepNext')))
        self.assertEqual(fmt.audit(path, True)['errors'], 0)
        self.assertEqual(len(pkg.doc.xpath('//w:br[@w:type="page"]', namespaces=fmt.NS)), 1)
        for style in pkg.styles.findall('w:style', fmt.NS):
            for flag in fmt.CONTROL_PROPERTIES:
                self.assertIsNone(style.find('w:pPr/w:' + flag, fmt.NS))

    def test_cleanup_preserves_content_and_clears_direct_and_inherited_controls(self):
        from docx import Document
        from docx.oxml import OxmlElement
        doc = Document()
        normal = doc.styles['Normal']
        normal.paragraph_format.keep_with_next = True
        normal.paragraph_format.keep_together = True
        normal.paragraph_format.page_break_before = True
        doc.add_heading('一、标题', 1)
        doc.add_paragraph('这是足够清晰的普通正文，测试继承的分页设置。').paragraph_format.keep_with_next = True
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = '表内文字'
        doc.add_paragraph('表1 示例', 'Caption')
        doc.add_paragraph('[1] 作者. 文献标题. 2026年。')
        doc.add_page_break()
        doc.sections[0].header.paragraphs[0].text = '页眉文字'
        doc.sections[0].header.paragraphs[0].paragraph_format.keep_with_next = True
        # Embedded field must survive the narrow XML repair.
        p = doc.add_paragraph('域内容所在正文')
        r = p.add_run()
        field = OxmlElement('w:fldChar')
        field.set(fmt.tag('fldCharType'), 'begin')
        r._r.append(field)
        path = self.path / 'inherited.docx'
        doc.save(path)
        before = fmt.Package(path)
        self.assertTrue(fmt.enabled(before.value(before.paragraphs[1], 'keepNext')))
        self.assertTrue(any(x['code'] == 'paragraph_pagination' for x in fmt.audit(path)['issues']))
        target = self.path / 'repaired.docx'
        report = fmt.clean_format(path, target)
        after = fmt.Package(target)
        self.assertFalse(fmt.enabled(after.value(after.paragraphs[1], 'keepNext')))
        self.assertFalse(fmt.enabled(after.value(after.paragraphs[1], 'keepLines')))
        self.assertFalse(fmt.enabled(after.value(after.paragraphs[0], 'keepNext')))
        self.assertEqual(after.style_id(after.paragraphs[0]), before.style_id(before.paragraphs[0]))
        self.assertEqual(after.value(after.paragraphs[0], 'outlineLvl'), before.value(before.paragraphs[0], 'outlineLvl'))
        self.assertEqual(fmt.audit(target, True)['errors'], 0)
        self.assertEqual(before.doc.xpath('//w:t/text()', namespaces=fmt.NS), after.doc.xpath('//w:t/text()', namespaces=fmt.NS))
        self.assertEqual(before.doc.xpath('//w:instrText/text()', namespaces=fmt.NS), after.doc.xpath('//w:instrText/text()', namespaces=fmt.NS))
        self.assertEqual(before.doc.xpath('//w:br/@w:type', namespaces=fmt.NS), after.doc.xpath('//w:br/@w:type', namespaces=fmt.NS))
        for name, part in before.parts.items():
            if name not in report['changed_parts']:
                self.assertEqual(part, after.parts[name])
        for index in range(len(before.paragraphs)):
            old, new = deepcopy(before.paragraphs[index]), deepcopy(after.paragraphs[index])
            fmt.strip_control_properties(old)
            self.assertEqual(E.tostring(old), E.tostring(new))
        with self.assertRaises(FileExistsError):
            fmt.clean_format(path, target)
        second = self.path / 'second.docx'
        fmt.clean_format(target, second)
        self.assertEqual(fmt.Package(target).parts, fmt.Package(second).parts)

    def test_style_bullets_and_cancellation(self):
        from docx import Document
        doc = Document()
        p = doc.add_paragraph('项目一', 'List Bullet')
        p2 = doc.add_paragraph('取消继承项目符号', 'List Bullet')
        fmt.sub(fmt.sub(p2._p.get_or_add_pPr(), 'numPr'), 'numId', val='0')
        path = self.path / 'bullets.docx'
        doc.save(path)
        pkg = fmt.Package(path)
        self.assertEqual(pkg.number_format(pkg.paragraphs[0]), 'bullet')
        self.assertIsNone(pkg.number_format(pkg.paragraphs[1]))
        target = self.path / 'clean-bullets.docx'
        fmt.clean_format(path, target)
        result = fmt.Package(target)
        self.assertIsNone(result.number_format(result.paragraphs[0]))
        self.assertEqual(fmt.audit(target, True)['errors'], 0)
        self.assertEqual(pkg.doc.xpath('//w:t/text()', namespaces=fmt.NS), result.doc.xpath('//w:t/text()', namespaces=fmt.NS))
        self.assertNotEqual(result.style_id(result.paragraphs[0]), 'ListBullet')

    def test_direct_overrides_are_not_hidden_by_style(self):
        from docx import Document
        from docx.shared import Pt
        path = self.build_sample()
        doc = Document(path)
        p = doc.paragraphs[2]
        p.runs[0].font.size = Pt(9)
        p.runs[0].font.name = 'Arial'
        p.paragraph_format.keep_with_next = True
        mutated = self.path / 'mutated.docx'
        doc.save(mutated)
        codes = {x['code'] for x in fmt.audit(mutated)['issues'] if x.get('paragraph') == 2}
        self.assertTrue({'font_size', 'latin_font', 'paragraph_pagination'}.issubset(codes), codes)

    def test_automatic_reference_numbers_are_not_silently_lost(self):
        from docx import Document
        doc = Document()
        doc.add_paragraph('作者. 文献标题. 2026。', 'List Number')
        source, target = self.path / 'references.docx', self.path / 'clean.docx'
        doc.save(source)
        with self.assertRaisesRegex(ValueError, '先将可见自动编号'):
            fmt.clean_format(source, target)
        self.assertFalse(target.exists())
        doc.paragraphs[0].text = '[1] 作者. 文献标题. 2026。'
        doc.paragraphs[0].style = doc.styles['Normal']
        doc.save(source)
        fmt.clean_format(source, target)
        self.assertEqual(fmt.audit(target, True)['errors'], 0)
        self.assertEqual(fmt.Package(source).doc.xpath('//w:t/text()', namespaces=fmt.NS),
                         fmt.Package(target).doc.xpath('//w:t/text()', namespaces=fmt.NS))

    def test_literal_symbols_are_not_deleted_as_formatting(self):
        from docx import Document
        doc = Document()
        doc.add_paragraph('• 原文实际字符')
        source, target = self.path / 'literal.docx', self.path / 'clean.docx'
        doc.save(source)
        with self.assertRaisesRegex(ValueError, '不能在纯格式清理时擅自删字'):
            fmt.clean_format(source, target)
        self.assertFalse(target.exists())

    def test_cleanup_preserves_native_notes_noteref_bookmarks_and_note_settings(self):
        from docx import Document
        from docx.oxml import OxmlElement
        doc = Document()
        doc.styles['Normal'].paragraph_format.keep_with_next = True
        p = doc.add_paragraph('原文与真实脚注')
        foot = OxmlElement('w:footnoteReference')
        foot.set(fmt.tag('id'), '2')
        p.add_run()._r.append(foot)
        p = doc.add_paragraph('首次尾注引用')
        start = OxmlElement('w:bookmarkStart')
        start.set(fmt.tag('id'), '42')
        start.set(fmt.tag('name'), 'sm_ref_0001')
        p._p.append(start)
        endnote = OxmlElement('w:endnoteReference')
        endnote.set(fmt.tag('id'), '2')
        p.add_run()._r.append(endnote)
        end = OxmlElement('w:bookmarkEnd')
        end.set(fmt.tag('id'), '42')
        p._p.append(end)
        p = doc.add_paragraph('重复引用')
        for kind, value in [('fldChar', 'begin'), ('instrText', ' NOTEREF sm_'),
                            ('instrText', 'ref_0001 '), ('fldChar', 'separate'),
                            ('t', '1'), ('fldChar', 'end')]:
            child = OxmlElement('w:' + kind)
            if kind == 'fldChar':
                child.set(fmt.tag('fldCharType'), value)
            else:
                child.text = value
            p.add_run()._r.append(child)
        source, target = self.path / 'native-notes.docx', self.path / 'clean-notes.docx'
        doc.save(source)
        with ZipFile(source) as z:
            parts = {name: z.read(name) for name in z.namelist()}
        ct = E.fromstring(parts['[Content_Types].xml'])
        rel = E.fromstring(parts['word/_rels/document.xml.rels'])
        for kind in ('footnote', 'endnote'):
            root = E.Element(fmt.tag(kind + 's'), nsmap={'w': fmt.W})
            for note_id, note_type in [('-1', 'separator'), ('0', 'continuationSeparator'), ('2', None)]:
                note = E.SubElement(root, fmt.tag(kind))
                note.set(fmt.tag('id'), note_id)
                if note_type:
                    note.set(fmt.tag('type'), note_type)
                paragraph = E.SubElement(note, fmt.tag('p'))
                if note_type is None:
                    pp = E.SubElement(paragraph, fmt.tag('pPr'))
                    fmt.sub(pp, 'keepNext', val='1')
                    fmt.sub(pp, 'ind', left='420', hanging='420')
                    run = E.SubElement(paragraph, fmt.tag('r'))
                    E.SubElement(run, fmt.tag(kind + 'Ref'))
                    E.SubElement(run, fmt.tag('tab'))
                    E.SubElement(run, fmt.tag('t')).text = '示例文献及原有说明'
            parts['word/' + kind + 's.xml'] = E.tostring(root)
            E.SubElement(ct, '{http://schemas.openxmlformats.org/package/2006/content-types}Override',
                         PartName='/word/' + kind + 's.xml',
                         ContentType='application/vnd.openxmlformats-officedocument.wordprocessingml.' + kind + 's+xml')
            E.SubElement(rel, '{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',
                         Id='rIdFixture' + kind, Target=kind + 's.xml',
                         Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/' + kind + 's')
        settings = E.fromstring(parts['word/settings.xml'])
        edn = fmt.sub(settings, 'endnotePr')
        fmt.sub(edn, 'pos', val='sectEnd')
        fmt.sub(edn, 'numFmt', val='decimal')
        fmt.sub(edn, 'numRestart', val='continuous')
        parts['word/settings.xml'] = E.tostring(settings)
        parts['[Content_Types].xml'] = E.tostring(ct)
        parts['word/_rels/document.xml.rels'] = E.tostring(rel)
        with ZipFile(source, 'w', ZIP_DEFLATED) as z:
            for name, data in parts.items():
                z.writestr(name, data)
        fmt.clean_format(source, target)
        after = fmt.Package(target)
        for part in ('word/document.xml', 'word/footnotes.xml', 'word/endnotes.xml'):
            expected = E.fromstring(parts[part])
            fmt.strip_control_properties(expected)
            actual = E.fromstring(after.parts[part])
            self.assertEqual(E.tostring(expected), E.tostring(actual), part)
        for part in ('word/settings.xml', 'word/_rels/document.xml.rels', '[Content_Types].xml'):
            self.assertEqual(parts[part], after.parts[part], part)
        self.assertEqual(fmt.audit(target, True)['errors'], 0)


if __name__ == '__main__':
    unittest.main()
