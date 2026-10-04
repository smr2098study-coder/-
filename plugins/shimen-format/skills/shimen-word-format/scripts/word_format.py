"""Build and audit DOCX with the user's Chinese academic formatting rules."""
import argparse
from collections import Counter
import json
from pathlib import Path
import re
from tempfile import TemporaryDirectory
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree as E

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
NS = {'w': W, 'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
PROFILES = {
    'body': ('宋体', 12, False, 200, 360),
    'title': ('宋体', 16, True, 0, 360),
    'h1': ('宋体', 16, True, 0, 360),
    'h2': ('宋体', 15, True, 0, 360),
    'h3': ('宋体', 14, True, 200, 360),
    'h4': ('宋体', 12, True, 200, 360),
    'h5': ('宋体', 12, False, 200, 360),
    'table': ('黑体', 10.5, False, 0, 276),
    'caption': ('黑体', 10.5, False, 0, 276),
}
PATTERNS = [r'^[一二三四五六七八九十百零]+、', r'^（[一二三四五六七八九十百零]+）',
            r'^\d+、', r'^（\d+）', r'^[①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳]']
FLAGS = ('keepNext', 'keepLines', 'pageBreakBefore')
CONTROL_PROPERTIES = FLAGS + ('numPr',)


def list_style_name(name):
    value = re.sub(r'[\s_-]', '', name).lower()
    return bool(re.fullmatch(r'list(?:bullet|number|paragraph|continue)?\d*', value)
                or re.match(r'^(列表|项目符号|编号列表)', value))


def literal_bullet(text):
    return text.lstrip().startswith(('•', '●', '▪', '■', '·')) or bool(re.match(r'^\s*[-*]\s+', text))


def strip_control_properties(root):
    count = Counter()
    for pp in root.findall('.//w:pPr', NS):
        for child in list(pp):
            name = E.QName(child).localname
            if name in CONTROL_PROPERTIES:
                pp.remove(child)
                count[name] += 1
    return count
P_ORDER = ('pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers '
           'pBdr shd tabs suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE '
           'autoSpaceDN bidi adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents '
           'suppressOverlap jc textDirection textAlignment textboxTightWrap outlineLvl divId cnfStyle '
           'rPr sectPr pPrChange').split()
R_ORDER = ('rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint '
           'noProof snapToGrid vanish webHidden color spacing w kern position sz szCs highlight u '
           'effect bdr shd fitText vertAlign rtl cs em lang eastAsianLayout specVanish oMath rPrChange').split()


def canonical_properties(root):
    """Keep core OOXML property order; preserve unknown extension order at the end."""
    for node, order in (('pPr', P_ORDER), ('rPr', R_ORDER)):
        properties = ([root] if E.QName(root).localname == node else []) + root.findall('.//w:' + node, NS)
        ranks = {tag(name): i for i, name in enumerate(order)}
        for prop in properties:
            prop[:] = sorted(prop, key=lambda child: ranks.get(child.tag, len(ranks)))


def tag(name):
    return '{' + W + '}' + name


def sub(parent, name, **attrs):
    el = parent.find(tag(name))
    if el is None:
        el = E.SubElement(parent, tag(name))
    for key, value in attrs.items():
        el.set(tag(key), str(value))
    return el


def enabled(value):
    return value is not None and value not in ('0', 'false', 'off')


def no_overwrite(source, target):
    target = Path(target)
    if source is not None and Path(source).resolve() == target.resolve():
        raise ValueError('输入输出不能是同一文件')
    if target.exists():
        raise FileExistsError('输出已存在，请使用新路径: ' + str(target))
    target.parent.mkdir(parents=True, exist_ok=True)
    return target


class Package:
    def __init__(self, path):
        with ZipFile(path) as z:
            self.parts = {name: z.read(name) for name in z.namelist()}
        self.doc = E.fromstring(self.parts['word/document.xml'])
        self.styles = E.fromstring(self.parts['word/styles.xml'])
        self.by_id = {s.get(tag('styleId')): s for s in self.styles.findall('w:style', NS)}
        self.default = next((sid for sid, s in self.by_id.items()
                             if s.get(tag('default')) == '1' and s.get(tag('type')) == 'paragraph'), 'Normal')
        self.defaults = self.styles.find('w:docDefaults', NS)
        self.numbering = E.fromstring(self.parts['word/numbering.xml']) if 'word/numbering.xml' in self.parts else None
        self.theme = E.fromstring(self.parts['word/theme/theme1.xml']) if 'word/theme/theme1.xml' in self.parts else None
        self.paragraphs = self.doc.xpath('//w:body//w:p', namespaces=NS)

    def chain(self, sid):
        chain, seen = [], set()
        while sid in self.by_id and sid not in seen:
            seen.add(sid)
            s = self.by_id[sid]
            chain.insert(0, s)
            base = s.find('w:basedOn', NS)
            sid = base.get(tag('val')) if base is not None else None
        return chain

    def style_id(self, p):
        s = p.find('w:pPr/w:pStyle', NS)
        return s.get(tag('val')) if s is not None else self.default

    def property_layers(self, p, run=None):
        kind = 'rPr' if run is not None else 'pPr'
        layers = []
        if self.defaults is not None:
            layers.append(self.defaults.find('w:' + kind + 'Default/w:' + kind, NS))
        layers.extend(s.find('w:' + kind, NS) for s in self.chain(self.style_id(p)))
        if run is None:
            layers.append(p.find('w:pPr', NS))
        else:
            rs = run.find('w:rPr/w:rStyle', NS)
            if rs is not None:
                layers.extend(s.find('w:rPr', NS) for s in self.chain(rs.get(tag('val'))))
            layers.append(run.find('w:rPr', NS))
        return [x for x in layers if x is not None]

    def value(self, p, name, attr='val', run=None):
        value = None
        for layer in self.property_layers(p, run):
            el = layer.find('w:' + name, NS)
            if el is not None:
                if el.get(tag(attr)) is not None:
                    value = el.get(tag(attr))
                elif attr == 'val':
                    value = '1'
        return value

    def bold(self, p, run):
        # b is an OOXML toggle in styles, but absolute in direct run formatting.
        value = False
        for layer in self.property_layers(p, run):
            el = layer.find('w:b', NS)
            if el is None:
                continue
            state = enabled(el.get(tag('val'), '1'))
            if layer.getparent() is not None and E.QName(layer.getparent()).localname == 'r':
                value = state
            elif state:
                value = not value
        return value

    def font(self, p, run, channel):
        # At each cascade level theme font takes precedence over literal font.
        theme_attr = {'eastAsia': 'eastAsiaTheme', 'ascii': 'asciiTheme', 'hAnsi': 'hAnsiTheme'}[channel]
        choice = None
        for layer in self.property_layers(p, run):
            fonts = layer.find('w:rFonts', NS)
            if fonts is None:
                continue
            if fonts.get(tag(theme_attr)) is not None:
                choice = ('theme', fonts.get(tag(theme_attr)))
            elif fonts.get(tag(channel)) is not None:
                choice = ('literal', fonts.get(tag(channel)))
        if choice is None:
            return None
        if choice[0] == 'literal':
            return choice[1]
        if self.theme is None:
            return 'unresolved-theme:' + choice[1]
        group = 'majorFont' if choice[1].startswith('major') else 'minorFont'
        base = self.theme.find('.//a:fontScheme/a:' + group, NS)
        if base is None:
            return 'unresolved-theme:' + choice[1]
        node = base.find('a:ea' if channel == 'eastAsia' else 'a:latin', NS)
        if node is not None and node.get('typeface'):
            return node.get('typeface')
        if channel == 'eastAsia':
            node = next((x for x in base.findall('a:font', NS) if x.get('script') == 'Hans'), None)
            if node is not None:
                return node.get('typeface')
        return 'unresolved-theme:' + choice[1]

    def number_format(self, p):
        num_id = self.value(p, 'numPr/w:numId')
        if num_id is None or num_id == '0':
            return None
        if self.numbering is None:
            return 'unresolved'
        level = self.value(p, 'numPr/w:ilvl') or '0'
        num = next((x for x in self.numbering.findall('w:num', NS) if x.get(tag('numId')) == num_id), None)
        if num is None:
            return 'unresolved'
        override = next((x for x in num.findall('w:lvlOverride', NS) if x.get(tag('ilvl')) == level), None)
        lv = override.find('w:lvl', NS) if override is not None else None
        if lv is None:
            aid = num.find('w:abstractNumId', NS)
            abstract = next((x for x in self.numbering.findall('w:abstractNum', NS)
                             if aid is not None and x.get(tag('abstractNumId')) == aid.get(tag('val'))), None)
            if abstract is not None:
                lv = next((x for x in abstract.findall('w:lvl', NS) if x.get(tag('ilvl')) == level), None)
        fmt = lv.find('w:numFmt', NS) if lv is not None else None
        return fmt.get(tag('val')) if fmt is not None else 'unresolved'

    def role(self, p):
        text = ''.join(p.xpath('.//w:t/text()', namespaces=NS)).strip()
        sid = self.style_id(p)
        s = self.by_id.get(sid)
        name_el = s.find('w:name', NS) if s is not None else None
        name = name_el.get(tag('val'), '') if name_el is not None else ''
        lower = name.lower()
        if p.xpath('ancestor::w:tc', namespaces=NS):
            return 'table'
        if re.match(r'toc|目录', lower):
            return 'toc'
        if lower == 'title' or name == '标题':
            return 'title'
        if re.match(r'^[表图]\s*\d', text) or 'caption' in lower or name == '题注':
            return 'caption'
        if p.xpath('.//w:drawing|.//w:pict', namespaces=NS):
            return 'image'
        if not text:
            return 'empty'
        match = re.match(r'heading\s*([1-5])|标题\s*([1-5])', lower)
        if match:
            return 'h' + (match.group(1) or match.group(2))
        for level, pattern in enumerate(PATTERNS, 1):
            if len(text) <= 90 and re.match(pattern, text):
                return 'candidate-h' + str(level)
        if re.match(r'^第.+章', text) or '【写作方向】' in text or '【建议字数】' in text:
            return 'special'
        if self.value(p, 'jc') == 'center' and len(text) < 80:
            return 'special'
        return 'body'

    def save(self, target):
        self.parts['word/document.xml'] = E.tostring(self.doc, encoding='UTF-8', xml_declaration=True, standalone=True)
        with ZipFile(target, 'w', ZIP_DEFLATED) as z:
            for name, value in self.parts.items():
                z.writestr(name, value)


def format_audit(path):
    """Strict control-property audit across stories, defaults, and styles."""
    pkg = Package(path)
    issues, counts = [], Counter()
    for part, data in pkg.parts.items():
        if not part.startswith('word/') or not part.endswith('.xml'):
            continue
        root = E.fromstring(data)
        for index, p in enumerate(root.findall('.//w:p', NS)):
            text = ''.join(p.xpath('.//w:t/text()', namespaces=NS))
            if pkg.role(p) == 'body' and literal_bullet(text):
                issues.append({'part': part, 'paragraph': index, 'code': 'literal_bullet', 'severity': 'error',
                               'message': '正文存在文字项目符号；核对用途，纯格式清理不能擅自删字'})
        for pp in root.findall('.//w:pPr', NS):
            for child in pp:
                name = E.QName(child).localname
                if name in CONTROL_PROPERTIES:
                    counts[name] += 1
                    issues.append({'part': part, 'code': name, 'severity': 'error',
                                   'message': '严格清理要求中不能保留该段落属性，包括val=0节点'})
        for style in root.findall('.//w:style', NS):
            name = style.find('w:name', NS)
            if style.get(tag('type')) == 'paragraph' and name is not None and list_style_name(name.get(tag('val'), '')):
                issues.append({'part': part, 'code': 'list_style', 'severity': 'error',
                               'message': '存在列表段落样式，需转换为无编号的普通段落样式'})
    return {'scope': '全Word XML部件的分页/编号属性检查，不验证视觉布局或文本语义',
            'counts': {key: counts[key] for key in CONTROL_PROPERTIES},
            'errors': len(issues), 'warnings': 0, 'issues': issues}


def audit(path, format_only=False):
    hygiene = format_audit(path)
    if format_only:
        return hygiene
    pkg = Package(path)
    issues, stats = list(hygiene['issues']), Counter()
    def add(index, code, message, severity='error'):
        issues.append({'paragraph': index, 'code': code, 'severity': severity, 'message': message})
    for index, p in enumerate(pkg.paragraphs):
        text = ''.join(p.xpath('.//w:t/text()', namespaces=NS)).strip()
        role = pkg.role(p)
        stats[role] += 1
        for flag in FLAGS:
            if enabled(pkg.value(p, flag)):
                stats['effective_' + flag] += 1
                add(index, 'paragraph_pagination', flag + '在标题或正文中仍生效')
        numfmt = pkg.number_format(p)
        if numfmt:
            stats['numbered_paragraphs'] += 1
            add(index, 'list', '列表格式为' + numfmt + '，需按文字编号要求处理')
        if text.startswith(('•', '●', '▪', '■', '·')):
            add(index, 'literal_bullet', '段首实际黑点字符，纯格式清理不得擅自删文字', 'warning')
        if re.match(r'^#{1,6}\s', text) or '**' in text:
            add(index, 'markdown', '疑似Markdown残留', 'warning')
        if '【写作方向】' in text or '【建议字数】' in text:
            add(index, 'draft_note', '有写作指导标签，正式成稿需整理', 'warning')
        if role.startswith('candidate-h'):
            add(index, 'heading_style', '疑似标题尚未使用对应Heading样式', 'warning')
            role = role.removeprefix('candidate-')
        if role not in PROFILES or not text:
            if role == 'special':
                add(index, 'role_review', '封面、章节标题或指导段落需人工确认角色', 'warning')
            continue
        font, size, bold, indent, line = PROFILES[role]
        if role.startswith('h'):
            level = int(role[1])
            if not re.match(PATTERNS[level - 1], text) and numfmt is None:
                add(index, 'heading_number', '编号与师门标题级别不匹配；既有章节体系需人工映射', 'warning')
            if pkg.value(p, 'outlineLvl') != str(level - 1):
                add(index, 'outline_level', '标题缺少对应大纲级别')
        firstchars = pkg.value(p, 'ind', 'firstLineChars')
        twip = pkg.value(p, 'ind', 'firstLine') or '0'
        if firstchars is not None:
            ok = firstchars == str(indent)
        else:
            ok = twip == str(round(size * 40) if indent else 0)
        if not ok:
            add(index, 'indent', '首行缩进与角色要求不符')
        if enabled(pkg.value(p, 'ind', 'hanging')) or enabled(pkg.value(p, 'ind', 'hangingChars')):
            add(index, 'hanging_indent', '存在冲突悬挂缩进')
        # Body/table line spacing is source-mandated; heading/caption spacing is a default.
        if role in ('body', 'table'):
            if pkg.value(p, 'spacing', 'line') != str(line) or pkg.value(p, 'spacing', 'lineRule') not in (None, 'auto'):
                add(index, 'line_spacing', '行距不是规定的倍数行距')
        if enabled(pkg.value(p, 'snapToGrid')):
            add(index, 'grid', '段落与网格对齐可能覆盖倍数行距', 'warning')
        seen = set()
        for run in p.xpath('.//w:r', namespaces=NS):
            rt = ''.join(run.xpath('.//w:t/text()', namespaces=NS))
            if not rt.strip():
                continue
            checks = []
            actual_size = pkg.value(p, 'sz', run=run)
            if actual_size != str(round(size * 2)):
                checks.append(('font_size', '部分文字字号不符'))
            if re.search(r'[\u3400-\u9fff]', rt) and pkg.font(p, run, 'eastAsia') != font:
                checks.append(('chinese_font', '部分中文字体不符或无法解析'))
            if re.search(r'[A-Za-z0-9]', rt) and pkg.font(p, run, 'ascii') != 'Times New Roman':
                checks.append(('latin_font', '部分英文数字字体不符或无法解析'))
            if role.startswith('h') or role == 'title':
                if pkg.bold(p, run) != bold:
                    checks.append(('bold', '标题加粗设置不符'))
            for code, message in checks:
                if code not in seen:
                    seen.add(code)
                    add(index, code, message)
    for i, tbl in enumerate(pkg.doc.findall('.//w:tbl', NS)):
        jc = tbl.find('w:tblPr/w:jc', NS)
        if jc is None or jc.get(tag('val')) != 'center':
            issues.append({'table': i, 'code': 'table_alignment', 'severity': 'error', 'message': '表格未显式居中，需排除表样式继承后确认'})
        for h in tbl.findall('.//w:trHeight', NS):
            if h.get(tag('hRule')) == 'exact':
                issues.append({'table': i, 'code': 'fixed_height', 'severity': 'warning', 'message': '固定行高可能截断内容'})
    return {'scope': '主文档结构审计，未做视觉验收；自动角色必须复核', 'paragraphs': len(pkg.paragraphs),
            'tables': len(pkg.doc.findall('.//w:tbl', NS)), 'stats': dict(stats),
            'errors': sum(x['severity'] == 'error' for x in issues),
            'warnings': sum(x['severity'] == 'warning' for x in issues), 'issues': issues}


def style_profile(style, role):
    from docx.shared import Pt, RGBColor
    font, size, bold, indent, line = PROFILES[role]
    style.font.name = 'Times New Roman'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = sub(rpr, 'rFonts', ascii='Times New Roman', hAnsi='Times New Roman', eastAsia=font, cs='Times New Roman')
    for attr in list(fonts.attrib):
        if 'theme' in E.QName(attr).localname.lower():
            del fonts.attrib[attr]
    sub(rpr, 'szCs', val=round(size * 2))
    sub(rpr, 'bCs', val='1' if bold else '0')
    ppr = style.element.get_or_add_pPr()
    ind = sub(ppr, 'ind')
    ind.attrib.clear()
    ind.set(tag('firstLineChars'), str(indent))
    ind.set(tag('firstLine'), str(round(size * 40) if indent else 0))
    sub(ppr, 'spacing', line=line, lineRule='auto', before='120' if role.startswith('h') else '0', after='0')
    sub(ppr, 'snapToGrid', val='0')
    sub(ppr, 'widowControl', val='1')
    for child in list(ppr):
        if E.QName(child).localname in CONTROL_PROPERTIES:
            ppr.remove(child)
    sub(ppr, 'jc', val='center' if role in ('title', 'caption') else 'both' if role == 'body' else 'left')
    if role.startswith('h'):
        sub(ppr, 'outlineLvl', val=int(role[1]) - 1)


def build(source, target):
    from docx import Document
    from docx.shared import Cm
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    source = Path(source)
    data = json.loads(source.read_text(encoding='utf-8'))
    doc = Document()
    for section in doc.sections:
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Cm(2.54)
    names = {'body': 'Normal', 'title': 'Title', **{'h' + str(i): 'Heading ' + str(i) for i in range(1, 6)},
             'table': 'Shimen Table Text', 'caption': 'Caption'}
    for role, name in names.items():
        style = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        if role != 'body':
            style.base_style = doc.styles['Normal']
        style_profile(style, role)
        style.next_paragraph_style = doc.styles['Normal']
    if data.get('title'):
        doc.add_paragraph(data['title'], 'Title')
    for block in data.get('blocks', []):
        kind = block['type']
        if kind == 'heading':
            level = block['level']
            if not isinstance(level, int) or not 1 <= level <= 5:
                raise ValueError('标题级别必须在1到5之间')
            doc.add_paragraph(block['text'], 'Heading ' + str(level))
        elif kind == 'paragraph':
            if literal_bullet(block['text']):
                raise ValueError('新建正文不得使用项目符号或无序列表；请使用普通段落文本')
            doc.add_paragraph(block['text'], 'Normal')
        elif kind == 'page_break':
            doc.add_page_break()
        elif kind == 'table':
            rows = block['rows']
            if not rows or not rows[0] or any(len(r) != len(rows[0]) for r in rows):
                raise ValueError('表格行不能为空且列数必须一致')
            if not block.get('caption'):
                raise ValueError('表格必须提供caption')
            doc.add_paragraph(block['caption'], 'Caption')
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = 'Table Grid'
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for row, values in zip(table.rows, rows):
                for cell, value in zip(row.cells, values):
                    cell.text = str(value)
                    cell.paragraphs[0].style = doc.styles['Shimen Table Text']
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        elif kind == 'image':
            if not block.get('caption'):
                raise ValueError('图片必须提供caption')
            width = block.get('width_cm', 12)
            if not 0 < width <= 15.92:
                raise ValueError('图片宽度必须在0到15.92cm之间')
            p = doc.add_paragraph()
            p.paragraph_format.alignment = 1
            p.paragraph_format.first_line_indent = Cm(0)
            sub(p._p.get_or_add_pPr(), 'ind', firstLineChars=0, firstLine=0)
            p.add_run().add_picture(str(source.parent / block['path']), width=Cm(width))
            doc.add_paragraph(block['caption'], 'Caption')
        else:
            raise ValueError('不支持的块类型: ' + kind)
    target = no_overwrite(source, target)
    canonical_properties(doc.element)
    canonical_properties(doc.styles.element)
    # python-docx templates also contain stylesWithEffects.xml, which is not
    # exposed by Document.styles. Run the package-wide cleaner before delivery.
    with TemporaryDirectory(prefix='shimen-build-') as temporary:
        intermediate = Path(temporary) / 'draft.docx'
        doc.save(intermediate)
        clean_format(intermediate, target)


def normalize_list_styles(styles, roots):
    """Rename de-numbered list styles, preserving all unrelated style properties."""
    existing = {s.get(tag('styleId')) for s in styles.findall('w:style', NS)}
    mapping = {}
    for style in styles.findall('w:style', NS):
        name = style.find('w:name', NS)
        if style.get(tag('type')) != 'paragraph' or name is None or not list_style_name(name.get(tag('val'), '')):
            continue
        old = style.get(tag('styleId'))
        new = 'ShimenPlain' + str(len(mapping) + 1)
        while new in existing:
            new += 'X'
        existing.add(new)
        mapping[old] = new
        style.set(tag('styleId'), new)
        name.set(tag('val'), 'Shimen Plain Paragraph ' + str(len(mapping)))
    for root in roots:
        for node in root.iter():
            if node.tag in {tag(x) for x in ('pStyle', 'basedOn', 'next', 'link', 'styleLink', 'numStyleLink')}:
                old = node.get(tag('val'))
                if old in mapping:
                    node.set(tag('val'), mapping[old])
    return mapping


def clean_format(source, target):
    """Clear control properties without rebuilding text, tables, fields, or media."""
    pkg = Package(source)
    target = no_overwrite(source, target)
    roots = {}
    for part, data in pkg.parts.items():
        if part.startswith('word/') and part.endswith('.xml'):
            roots[part] = E.fromstring(data)
    # Numeric labels are visible content: do not silently delete auto-numbered
    # headings or references. Convert them to verified literal labels first.
    for part, root in roots.items():
        for index, p in enumerate(root.findall('.//w:p', NS)):
            numfmt = pkg.number_format(p)
            text = ''.join(p.xpath('.//w:t/text()', namespaces=NS)).strip()
            if pkg.role(p) == 'body' and literal_bullet(text):
                raise ValueError('段首文字符号属于内容，先核对其用途，不能在纯格式清理时擅自删字: '
                                 + part + ' paragraph ' + str(index))
            if numfmt not in (None, 'bullet', 'none'):
                raise ValueError('先将可见自动编号准确转换为普通文字以保留标题/参考文献: '
                                 + part + ' paragraph ' + str(index))
    snapshots = {part: E.tostring(root) for part, root in roots.items()}
    removed = Counter()
    for root in roots.values():
        removed.update(strip_control_properties(root))
    mapping = normalize_list_styles(roots['word/styles.xml'], list(roots.values()))
    # Shadow style parts, if present, must use the same renamed identifiers.
    for part, root in roots.items():
        if part != 'word/styles.xml':
            for s in root.findall('.//w:style', NS):
                old = s.get(tag('styleId'))
                if old in mapping:
                    s.set(tag('styleId'), mapping[old])
                    name = s.find('w:name', NS)
                    if name is not None:
                        name.set(tag('val'), 'Shimen Plain Paragraph ' + str(list(mapping).index(old) + 1))
    changed = []
    for part, root in roots.items():
        if E.tostring(root) != snapshots[part]:
            pkg.parts[part] = E.tostring(root, encoding='UTF-8', xml_declaration=True, standalone=True)
            changed.append(part)
    with ZipFile(target, 'w', ZIP_DEFLATED) as z:
        for part, data in pkg.parts.items():
            z.writestr(part, data)
    verification = format_audit(target)
    if verification['errors']:
        raise ValueError('格式清理后检查未通过，请查看输出并修复')
    return {'removed': {name: removed[name] for name in CONTROL_PROPERTIES},
            'renamed_list_styles': len(mapping), 'changed_parts': changed,
            'verification': verification['counts'], 'scope': '仅分页/编号及列表样式，保留文字内容和其他格式'}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest='command', required=True)
    for command in ('build', 'audit', 'clean-format'):
        p = commands.add_parser(command)
        p.add_argument('input')
        p.add_argument('--out', required=True)
        if command == 'audit':
            p.add_argument('--format-only', action='store_true')
    args = parser.parse_args()
    if args.command == 'build':
        build(args.input, args.out)
        print('Created:', args.out)
    elif args.command == 'audit':
        result = audit(args.input, args.format_only)
        target = no_overwrite(args.input, args.out)
        target.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
        print(json.dumps({k: v for k, v in result.items() if k != 'issues'}, ensure_ascii=False))
        raise SystemExit(1 if result['errors'] else 0)
    else:
        print(json.dumps(clean_format(args.input, args.out), ensure_ascii=False))


if __name__ == '__main__':
    main()
